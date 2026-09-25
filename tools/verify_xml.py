"""Independent second check of an XML made by eSAD Item Builder.

It re-reads the Excel packing list with openpyxl (the app uses SheetJS), and checks the
XML against it using only the audit report's line-to-item mapping and answers.

    python3 tools/verify_xml.py PL.xlsx out.xml out-audit.csv [original.xml]

Exit code 0 means every check passed.
"""
import csv
import json
import re
import sys
import xml.etree.ElementTree as ET
from decimal import Decimal

import openpyxl

HERE = __file__.rsplit("/", 1)[0]


def cents(v):
    return int((Decimal(str(v)) * 100).to_integral_value())


def grams(v):
    return int((Decimal(str(v)) * 1000).to_integral_value())


def number(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def largest_remainder(total, weights):
    W = sum(weights)
    w = weights if W > 0 else [1] * len(weights)
    WW = W if W > 0 else len(weights)
    raw = [total * x / WW for x in w]
    base = [int(x // 1) for x in raw]
    left = total - sum(base)
    for _, i in sorted(((raw[i] - base[i], i) for i in range(len(w))), key=lambda t: (-t[0], t[1])):
        if left <= 0:
            break
        base[i] += 1
        left -= 1
    return base


def read_audit(path):
    rows = list(csv.reader(open(path, encoding="utf-8-sig")))
    answers, mapping, section = {}, {}, None
    for r in rows:
        if not r or not any(r):
            section = None
            continue
        if r[0] == "Excel problems and answers":
            section = "answers"
            continue
        if r[0] == "Excel line" and len(r) > 7:
            section = "lines"
            continue
        if section == "answers" and len(r) >= 3:
            answers[r[0]] = r[2]
        elif section == "lines" and len(r) >= 8:
            mapping[int(float(r[0]))] = r[7]
    return answers, mapping


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    xlsx, out_xml, audit = sys.argv[1:4]
    original = sys.argv[4] if len(sys.argv) > 4 else None
    tar_js = open(__import__("os").environ.get("TARIFF_JS", f"{HERE}/../tariff-data.js"), encoding="utf-8").read()
    tariff = json.loads(tar_js[tar_js.index("{"):tar_js.rindex("}") + 1])
    t8 = {l["c"]: l for l in tariff["lines"]}
    answers, mapping = read_audit(audit)
    results = []

    def check(ok, title, detail=""):
        results.append((ok, title, detail))

    # ---- Excel, read independently ----
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    hdr = next(r for r in range(1, 30) if any(re.search(r"HS\s*CODE", str(c.value or ""), re.I) for c in ws[r]))
    heads = {c.column: re.sub(r"\s+", " ", str(c.value or "")) for c in ws[hdr]}
    col = lambda rx, exclude=(): next(k for k, v in heads.items() if re.search(rx, v, re.I) and k not in exclude)
    c_no, c_amt, c_net, c_gross, c_pk = col(r"^\s*(序号|No\.?)"), col(r"AMOUNT"), col(r"Total\s*N\.?\s*W|总净重"), col(r"G\.?\s*W|毛重"), col(r"PKGS|箱数")
    merges = [m for m in ws.merged_cells.ranges]
    top_of = lambda r, c: next(((m.min_row, m.min_col) for m in merges if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), (r, c))
    lines = []
    for r in range(hdr + 1, ws.max_row + 1):
        no = ws.cell(r, c_no).value
        if isinstance(no, str) and re.search(r"total|合计", no, re.I):
            break
        if not isinstance(no, (int, float)):
            continue
        ref = lambda c: f"{openpyxl.utils.get_column_letter(c)}{r}"
        amt = answers.get(ref(c_amt))
        amt = number(amt.split("answer: ")[-1]) if amt and "answer:" in amt else number(ws.cell(r, c_amt).value)
        net = answers.get(ref(c_net))
        net = number(net.split("answer: ")[-1]) if net and "answer:" in net else number(ws.cell(r, c_net).value)
        gr, gc = top_of(r, c_gross)
        pr, pc = top_of(r, c_pk)
        gcell = f"{openpyxl.utils.get_column_letter(gc)}{gr}"
        g_ans = answers.get(ref(c_gross), "")
        pk = number(ws.cell(r, c_pk).value) if (pr, pc) == (r, c_pk) else 0
        lines.append({"no": int(no), "row": r, "amt": cents(amt) if amt is not None else None, "net": grams(net) if net is not None else None,
                      "group": gcell, "g_top": (gr, gc) == (r, c_gross), "g_ans": g_ans,
                      "gross": number(ws.cell(r, c_gross).value), "pk": int(pk or 0)})
    check(len(lines) == len(mapping), "Excel line count matches the audit report", f"{len(lines)} lines in Excel, {len(mapping)} in the audit")
    check(all(l["amt"] is not None and l["net"] is not None for l in lines), "Every Excel line has a price and net weight (from Excel or an answer)")

    # carton groups and gross shares, recomputed here
    by_no = {l["no"]: l for l in lines}
    for l in lines:
        m = re.search(r"same carton as line (\d+)", l["g_ans"])
        if m:  # the whole merged cell joins the other carton
            old, new = l["group"], by_no[int(m.group(1))]["group"]
            l["g_top"] = False
            for x in lines:
                if x["group"] == old:
                    x["group"] = new
        elif "answer:" in l["g_ans"]:
            l["gross"] = number(l["g_ans"].split("answer: ")[-1])
    groups = {}
    for l in lines:
        groups.setdefault(l["group"], []).append(l)
    share = {}
    for g in groups.values():
        total = sum(grams(l["gross"]) for l in g if l["g_top"] and l["gross"] is not None)
        w = [l["net"] or 0 for l in g]
        parts = [x * 10 for x in largest_remainder(total // 10, w)] if total % 10 == 0 else largest_remainder(total, w)
        for l, p in zip(g, parts):
            share[l["no"]] = p

    # ---- XML ----
    raw = open(out_xml, encoding="utf-8").read()
    root = ET.fromstring(raw.encode("utf-8"))
    items = root.findall("Item")
    by_rank = {i.findtext("Rank"): i for i in items}
    new = {}
    for no, place in mapping.items():
        m = re.match(r"new item (\d+)", place)
        if m:
            new.setdefault(m.group(1), []).append(by_no[no])
    unplaced = [no for no, p in mapping.items() if p == "NOT PLACED"]
    check(not unplaced, "No Excel line is left unplaced", f"lines {unplaced}" if unplaced else "")
    bad = []
    for rank, ls in new.items():
        it = by_rank.get(rank)
        if it is None:
            bad.append(f"item {rank} missing from XML")
            continue
        g = lambda p: (it.findtext(p) or "").strip()
        if cents(g("Tarification/Item_price")) != sum(l["amt"] for l in ls):
            bad.append(f"item {rank}: price {g('Tarification/Item_price')} ≠ Excel {sum(l['amt'] for l in ls)/100:.2f}")
        if cents(g("Valuation_item/Item_Invoice/Amount_foreign_currency")) != sum(l["amt"] for l in ls):
            bad.append(f"item {rank}: invoice amount differs from Excel")
        if grams(g("Valuation_item/Weight_itm/Net_weight_itm")) != sum(l["net"] for l in ls):
            bad.append(f"item {rank}: net {g('Valuation_item/Weight_itm/Net_weight_itm')} ≠ Excel {sum(l['net'] for l in ls)/1000}")
        if grams(g("Valuation_item/Weight_itm/Gross_weight_itm")) != sum(share[l["no"]] for l in ls):
            bad.append(f"item {rank}: gross {g('Valuation_item/Weight_itm/Gross_weight_itm')} ≠ recomputed {sum(share[l['no']] for l in ls)/1000}")
        code = g("Tarification/HScode/Commodity_code")
        if code not in t8:
            bad.append(f"item {rank}: {code} not in tariff")
        else:
            want = [u for u in re.split(r"[\s,/]+", t8[code]["u"]) if u]
            have = [(s.findtext("Supplementary_unit_code") or "").strip() for s in it.findall("Tarification/Supplementary_unit")]
            if [h for h in have if h] != want:
                bad.append(f"item {rank}: units {have} but tariff says {want}")
            if g("Goods_description/Harmonized_System_Code_Description") != t8[code]["d"]:
                bad.append(f"item {rank}: tariff description differs")
    check(not bad, f"All {len(new)} new items match the Excel lines they came from", "; ".join(bad[:6]))
    new_lines = [l for ls in new.values() for l in ls]
    pk_new = sum(int(float(by_rank[r].findtext("Packages/Number_of_packages"))) for r in new)
    check(pk_new == sum(l["pk"] for l in new_lines), "Packages of the new items equal the cartons of their Excel lines", f"{pk_new} vs {sum(l['pk'] for l in new_lines)}")
    all_bad = [i.findtext("Rank") for i in items if (i.findtext("Tarification/HScode/Commodity_code") or "").strip() not in t8]
    check(not all_bad, "Every item's HS code is in the official tariff", f"items {all_bad}" if all_bad else f"{len(items)} items")
    check(int(root.findtext("Property/Nbers/Total_number_of_items")) == len(items), "Header item count is right")
    check(int(float(root.findtext("Property/Nbers/Total_number_of_packages"))) == sum(int(float(i.findtext("Packages/Number_of_packages"))) for i in items), "Header package count is right")

    if original:
        o = open(original, encoding="utf-8").read()
        o_items = [t for t in re.findall(r"<Item>[\s\S]*?</Item>", o) if "<Commodity_code/>" not in t]
        out_items = re.findall(r"<Item>[\s\S]*?</Item>", raw)
        changed = [re.search(r"<Rank>(\d+)</Rank>", t).group(1) for t in o_items if t not in out_items]
        check(not changed, "Items from the original XML are unchanged", f"changed: {changed} (only allowed for HS codes you corrected)" if changed else f"{len(o_items)} identical")
        o_head = o[:o.index("<Item>")]
        r_head = raw[:raw.index("<Item>")]
        diff = [a for a, b in zip(o_head.splitlines(), r_head.splitlines()) if a != b]
        check(all(re.search(r"Total_number_of_items|Total_number_of_packages|Gross_weight|Total_weight", d) for d in diff), "Header changed only in the totals", f"{len(diff)} header lines differ")

    ok = all(r[0] for r in results)
    for passed, title, detail in results:
        print(("PASS " if passed else "FAIL ") + title + (f"  ({detail})" if detail else ""))
    print("\nALL CHECKS PASSED" if ok else "\nSOME CHECKS FAILED")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
